"""Report export — Excel and PDF generation with batched row processing."""

from __future__ import annotations

import asyncio
import io
from collections.abc import AsyncIterator, Callable, Sequence
from datetime import UTC, datetime

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from webstudio_backend.infrastructure.repositories.report_filters import ExportFormat, ReportType
from webstudio_backend.infrastructure.repositories.report_repository import (
    AggregateReportRow,
    AuditReportRow,
    InventoryReportRow,
    NotificationReportRow,
    SalesExportRow,
)

PDF_ROWS_PER_PAGE = 40


def _format_storage(value: object | None, unit: str | None, storage_type: str | None) -> str:
    if value is None or str(value).strip() == "":
        return ""
    text = str(value).strip()
    if text.endswith(".00"):
        text = text[:-3]
    elif "." in text:
        text = text.rstrip("0").rstrip(".")
    parts = [text]
    if unit:
        parts.append(str(unit))
    joined = "".join(parts)
    if storage_type:
        joined = f"{joined} {storage_type}"
    return joined.strip()


def format_configuration(
    *,
    cpu: str | None,
    ram_gb: int | None,
    storage_value: object | None,
    storage_unit: str | None,
    storage_type: str | None,
    gpu: str | None = None,
    display: str | None = None,
) -> str:
    """Human-readable single-cell spec, e.g. 'i5-1240P / 16GB / 512GB SSD / 15.6\"'."""
    parts: list[str] = []
    if cpu:
        parts.append(str(cpu).strip())
    if ram_gb:
        parts.append(f"{ram_gb}GB")
    storage = _format_storage(storage_value, storage_unit, storage_type)
    if storage:
        parts.append(storage)
    if gpu:
        parts.append(str(gpu).strip())
    if display:
        parts.append(str(display).strip())
    return " / ".join(parts)


class ReportExportService:
    async def build_export(
        self,
        *,
        report_type: ReportType,
        export_format: ExportFormat,
        title: str,
        headers: Sequence[str],
        row_batches: AsyncIterator[list[Sequence[object]]],
    ) -> tuple[bytes, str, str]:
        collected = await self._collect_batches(row_batches)
        if export_format is ExportFormat.XLSX:
            content = await asyncio.to_thread(self._build_xlsx, title, headers, collected)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            extension = "xlsx"
        else:
            content = await asyncio.to_thread(self._build_pdf, title, headers, collected)
            media_type = "application/pdf"
            extension = "pdf"
        filename = (
            f"{report_type.value}-report-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}.{extension}"
        )
        return content, media_type, filename

    @staticmethod
    async def _collect_batches(
        batches: AsyncIterator[list[Sequence[object]]],
    ) -> list[list[Sequence[object]]]:
        collected: list[list[Sequence[object]]] = []
        async for batch in batches:
            collected.append(batch)
        return collected

    async def build_inventory_export(
        self,
        *,
        export_format: ExportFormat,
        rows: Sequence[InventoryReportRow],
    ) -> tuple[bytes, str, str]:
        """Stock export: a 'By Model' summary sheet + a per-serial 'Detail' sheet."""
        summary = self.inventory_summary_rows(rows)
        detail = [self.inventory_detail_row_values(row) for row in rows]
        sheets = [
            ("By Model", self.inventory_summary_headers(), summary),
            ("Detail", self.inventory_detail_headers(), detail),
        ]
        stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
        if export_format is ExportFormat.XLSX:
            content = await asyncio.to_thread(self._build_multi_sheet_xlsx, sheets)
            return (
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"inventory-report-{stamp}.xlsx",
            )
        content = await asyncio.to_thread(
            self._build_multi_table_pdf, "WEBSTUDIO IMS Stock Report", sheets
        )
        return content, "application/pdf", f"inventory-report-{stamp}.pdf"

    async def build_sales_export(
        self,
        *,
        export_format: ExportFormat,
        rows: Sequence[SalesExportRow],
        additional_by_guid: dict[str, str],
    ) -> tuple[bytes, str, str]:
        headers = self.sales_export_headers()
        data_rows = [
            self.sales_export_row_values(
                row, additional_by_guid.get(row.tally_voucher_guid or "", "")
            )
            for row in rows
        ]
        sheets = [("Sales", headers, data_rows)]
        stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
        if export_format is ExportFormat.XLSX:
            content = await asyncio.to_thread(self._build_multi_sheet_xlsx, sheets)
            return (
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"sales-report-{stamp}.xlsx",
            )
        content = await asyncio.to_thread(
            self._build_multi_table_pdf, "WEBSTUDIO IMS Sales Report", sheets
        )
        return content, "application/pdf", f"sales-report-{stamp}.pdf"

    def _build_multi_sheet_xlsx(
        self,
        sheets: Sequence[tuple[str, Sequence[str], Sequence[Sequence[object]]]],
    ) -> bytes:
        workbook = Workbook(write_only=True)
        for title, headers, data_rows in sheets:
            worksheet = workbook.create_sheet(title=title[:31])
            worksheet.append(list(headers))
            for row in data_rows:
                worksheet.append([self._cell_value(value) for value in row])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _build_multi_table_pdf(
        self,
        title: str,
        sheets: Sequence[tuple[str, Sequence[str], Sequence[Sequence[object]]]],
    ) -> bytes:
        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=24,
        )
        styles = getSampleStyleSheet()
        story: list[object] = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
        for sheet_title, headers, data_rows in sheets:
            story.append(Paragraph(sheet_title, styles["Heading2"]))
            story.append(Spacer(1, 6))
            rows: list[list[object]] = [list(headers)]
            rows.extend(list(row) for row in data_rows)
            for chunk_start in range(0, len(rows), PDF_ROWS_PER_PAGE):
                chunk = rows[chunk_start : chunk_start + PDF_ROWS_PER_PAGE]
                table = Table(
                    [[self._pdf_cell(value) for value in row] for row in chunk],
                    repeatRows=1 if chunk_start == 0 else 0,
                )
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                            ("FONTSIZE", (0, 0), (-1, -1), 8),
                            (
                                "ROWBACKGROUNDS",
                                (0, 1),
                                (-1, -1),
                                [colors.whitesmoke, colors.lightgrey],
                            ),
                        ],
                    ),
                )
                story.append(table)
                story.append(Spacer(1, 12))
        document.build(story)
        return buffer.getvalue()

    # ---- Inventory (stock) export -----------------------------------------

    def inventory_detail_headers(self) -> list[str]:
        return [
            "Serial Number",
            "Brand",
            "Model Number",
            "Model Name",
            "Configuration",
            "Color",
            "Location",
            "Status",
            "Purchase Date",
            "Purchase Price",
            "Archived",
        ]

    def inventory_detail_row_values(self, row: InventoryReportRow) -> list[object]:
        return [
            row.serial_number,
            row.brand_name,
            row.model_number,
            row.model_name,
            format_configuration(
                cpu=row.cpu,
                ram_gb=row.ram_gb,
                storage_value=row.storage_value,
                storage_unit=row.storage_unit,
                storage_type=row.storage_type,
                gpu=row.gpu,
                display=row.display,
            ),
            row.color,
            row.location_name,
            row.status,
            row.purchase_date.isoformat() if row.purchase_date else "",
            row.purchase_price if row.purchase_price is not None else "",
            row.is_archived,
        ]

    def inventory_summary_headers(self) -> list[str]:
        return [
            "Brand",
            "Model Number",
            "Model Name",
            "Configuration",
            "Quantity",
            "Serial Numbers",
        ]

    def inventory_summary_rows(self, rows: Sequence[InventoryReportRow]) -> list[list[object]]:
        """Group per model: brand, model, config, quantity and the serial list."""
        groups: dict[tuple, dict] = {}
        order: list[tuple] = []
        for row in rows:
            key = (
                row.brand_name,
                row.model_number,
                row.model_name,
                format_configuration(
                    cpu=row.cpu,
                    ram_gb=row.ram_gb,
                    storage_value=row.storage_value,
                    storage_unit=row.storage_unit,
                    storage_type=row.storage_type,
                    gpu=row.gpu,
                    display=row.display,
                ),
            )
            if key not in groups:
                groups[key] = {"serials": []}
                order.append(key)
            if row.serial_number:
                groups[key]["serials"].append(row.serial_number)
        summary: list[list[object]] = []
        for key in order:
            brand, model_number, model_name, config = key
            serials = groups[key]["serials"]
            summary.append(
                [
                    brand,
                    model_number,
                    model_name,
                    config,
                    len(serials),
                    ", ".join(serials),
                ]
            )
        return summary

    # ---- Sales export ------------------------------------------------------

    def sales_export_headers(self) -> list[str]:
        return [
            "Sold At",
            "Invoice",
            "Customer",
            "Brand",
            "Model Number",
            "Model Name",
            "Serial Number",
            "Configuration",
            "Color",
            "Quantity",
            "Amount",
            "Amount (excl. GST)",
            "CGST",
            "SGST",
            "IGST",
            "Cess",
            "Payment Mode",
            "Salesperson",
            "Source",
            "Additional Products",
        ]

    def sales_export_row_values(
        self, row: SalesExportRow, additional_products: str = ""
    ) -> list[object]:
        return [
            row.sold_at.isoformat() if row.sold_at else "",
            row.invoice_number,
            row.customer_name or "",
            row.brand_name,
            row.model_number,
            row.model_name,
            row.serial_number,
            format_configuration(
                cpu=row.cpu,
                ram_gb=row.ram_gb,
                storage_value=row.storage_value,
                storage_unit=row.storage_unit,
                storage_type=row.storage_type,
                gpu=row.gpu,
            ),
            row.color or "",
            1,
            row.sale_amount if row.sale_amount is not None else "",
            row.sale_amount_excluding_gst if row.sale_amount_excluding_gst is not None else "",
            row.sale_cgst_amount if row.sale_cgst_amount is not None else "",
            row.sale_sgst_amount if row.sale_sgst_amount is not None else "",
            row.sale_igst_amount if row.sale_igst_amount is not None else "",
            row.sale_cess_amount if row.sale_cess_amount is not None else "",
            row.payment_mode or "",
            row.recorded_by_display_name or "",
            row.sale_source,
            additional_products,
        ]

    def aggregate_headers(self) -> list[str]:
        return [
            "Group ID",
            "Group Name",
            "Available",
            "Sold",
            "Received",
            "Reserved",
            "Archived",
            "Total",
        ]

    def aggregate_row_values(self, row: AggregateReportRow) -> list[object]:
        return [
            row.group_id,
            row.group_name,
            row.available,
            row.sold,
            row.received,
            row.reserved,
            row.archived,
            row.total,
        ]

    def audit_headers(self) -> list[str]:
        return [
            "ID",
            "Entity Type",
            "Entity ID",
            "Action",
            "Source",
            "Actor",
            "Actor User ID",
            "Serial Number",
            "Description",
            "Created At",
        ]

    def audit_row_values(self, row: AuditReportRow) -> list[object]:
        return [
            str(row.id),
            row.entity_type,
            row.entity_id,
            row.action,
            row.source,
            row.actor_display_name or "",
            row.actor_user_id or "",
            row.serial_number or "",
            row.description or "",
            row.created_at.isoformat(),
        ]

    def notification_headers(self) -> list[str]:
        return [
            "ID",
            "Type",
            "Title",
            "Description",
            "Category",
            "Severity",
            "Status",
            "Created At",
            "Resolved At",
        ]

    def notification_row_values(self, row: NotificationReportRow) -> list[object]:
        return [
            row.id,
            row.notification_type,
            row.title,
            row.description,
            row.category,
            row.severity,
            row.status,
            row.created_at.isoformat(),
            row.resolved_at.isoformat() if row.resolved_at else "",
        ]

    def _build_xlsx(
        self,
        title: str,
        headers: Sequence[str],
        batches: list[list[Sequence[object]]],
    ) -> bytes:
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet(title=title[:31])
        worksheet.append(list(headers))
        for batch in batches:
            for row in batch:
                worksheet.append([self._cell_value(value) for value in row])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _build_pdf(
        self,
        title: str,
        headers: Sequence[str],
        batches: list[list[Sequence[object]]],
    ) -> bytes:
        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=24,
        )
        styles = getSampleStyleSheet()
        story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

        rows: list[list[object]] = [list(headers)]
        for batch in batches:
            rows.extend(batch)

        for chunk_start in range(0, len(rows), PDF_ROWS_PER_PAGE):
            chunk = rows[chunk_start : chunk_start + PDF_ROWS_PER_PAGE]
            table = Table(
                [[self._pdf_cell(value) for value in row] for row in chunk],
                repeatRows=1 if chunk_start == 0 else 0,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
                    ],
                ),
            )
            story.append(table)
            story.append(Spacer(1, 12))

        document.build(story)
        return buffer.getvalue()

    @staticmethod
    def _cell_value(value: object) -> str | int | float | bool:
        if value is None:
            return ""
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return value
        return str(value)

    @staticmethod
    def _pdf_cell(value: object) -> str:
        if value is None or isinstance(value, bool):
            return "" if value is None else str(value)
        return str(value)


def map_row_batches(
    batches: AsyncIterator[list[object]],
    mapper: Callable[[object], Sequence[object]],
) -> AsyncIterator[list[Sequence[object]]]:
    async def _generator() -> AsyncIterator[list[Sequence[object]]]:
        async for batch in batches:
            yield [mapper(row) for row in batch]

    return _generator()
