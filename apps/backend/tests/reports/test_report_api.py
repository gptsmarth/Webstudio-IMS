"""Report API, export, authorization, and performance tests."""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from webstudio_backend.infrastructure.audit.audit_actor import AuditActor
from webstudio_backend.infrastructure.database.enums import (
    InventoryStatus,
    LocationType,
    NotificationSeverity,
    NotificationType,
    StorageType,
    StorageUnit,
)
from webstudio_backend.infrastructure.repositories import (
    BrandRepository,
    InventoryItemRepository,
    LocationRepository,
    ProductModelRepository,
)
from webstudio_backend.infrastructure.repositories.report_filters import ReportFilters
from webstudio_backend.infrastructure.repositories.report_repository import ReportRepository, STREAM_BATCH_SIZE
from webstudio_backend.services.notification_service import NotificationService
from webstudio_backend.services.report_service import ReportService
from webstudio_backend.services.sale_service import SaleService


async def _seed_report_data(db_session: AsyncSession) -> tuple[int, int]:
    brand = await BrandRepository(db_session).create("ASUS")
    warehouse = await LocationRepository(db_session).create(
        "Warehouse",
        location_type=LocationType.WAREHOUSE,
    )
    store = await LocationRepository(db_session).create(
        "Store",
        location_type=LocationType.RETAIL_FLOOR,
    )
    product_model = await ProductModelRepository(db_session).create(
        brand_id=brand.id,
        model_number="X1502ZA",
        model_name="Vivobook 15",
        cpu="Intel i5",
        ram_gb=16,
        storage_value=Decimal("512"),
        storage_unit=StorageUnit.GB,
        storage_type=StorageType.SSD,
    )
    inventory_repo = InventoryItemRepository(db_session)
    available_item = await inventory_repo.create(
        serial_number="SN-RPT-001",
        product_model_id=product_model.id,
        color="Black",
        current_location_id=warehouse.id,
        status=InventoryStatus.AVAILABLE,
        purchase_date=date.today(),
    )
    sold_item = await inventory_repo.create(
        serial_number="SN-RPT-002",
        product_model_id=product_model.id,
        color="Silver",
        current_location_id=store.id,
        status=InventoryStatus.AVAILABLE,
    )
    actor = AuditActor(user_id=1, display_name="Main Admin", role="main_admin")
    await SaleService(db_session).reflect_manual_sale(
        inventory_item_id=sold_item.id,
        invoice_number="INV-RPT-001",
        customer_name="Acme Corp",
        payment_mode="Cash",
        sale_date=date.today(),
        remarks=None,
        actor=actor,
    )
    await NotificationService(db_session).create_notification(
        notification_type=NotificationType.INVENTORY_ALERT,
        severity=NotificationSeverity.WARNING,
        title="Low stock",
        message="Warehouse threshold reached",
    )
    await db_session.commit()
    return brand.id, warehouse.id


@pytest.mark.asyncio
async def test_inventory_report_requires_auth(api_client: AsyncClient) -> None:
    response = await api_client.get("/api/v1/reports/inventory")
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_inventory_report_generation(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    brand_id, location_id = await _seed_report_data(db_session)
    response = await api_client.get(
        "/api/v1/reports/inventory",
        params={"brand_id": brand_id, "location_id": location_id},
        headers=main_admin_headers,
    )
    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["report_type"] == "inventory"
    assert payload["summary"]["total_rows"] >= 1
    assert any(row["serial_number"] == "SN-RPT-001" for row in payload["rows"])


@pytest.mark.asyncio
async def test_sales_report_generation(
    api_client: AsyncClient,
    db_session: AsyncSession,
    admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get("/api/v1/reports/sales", headers=admin_headers)
    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["report_type"] == "sales"
    assert len(payload["rows"]) == 1
    assert payload["rows"][0]["invoice_number"] == "INV-RPT-001"


@pytest.mark.asyncio
async def test_audit_report_generation(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get("/api/v1/reports/audit", headers=main_admin_headers)
    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["report_type"] == "audit"
    assert len(payload["rows"]) >= 1


@pytest.mark.asyncio
async def test_notifications_report_generation(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get("/api/v1/reports/notifications", headers=main_admin_headers)
    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["report_type"] == "notification"
    assert payload["rows"][0]["title"] == "Low stock"


@pytest.mark.asyncio
async def test_salesperson_cannot_access_reports(
    api_client: AsyncClient,
    db_session: AsyncSession,
    salesperson_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get("/api/v1/reports/inventory", headers=salesperson_headers)
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_export_inventory_xlsx(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get(
        "/api/v1/reports/export",
        params={"report_type": "inventory", "format": "xlsx"},
        headers=main_admin_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Serial Number"
    assert any(row[0] == "SN-RPT-001" for row in rows[1:])


@pytest.mark.asyncio
async def test_export_sales_pdf(
    api_client: AsyncClient,
    db_session: AsyncSession,
    admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get(
        "/api/v1/reports/export",
        params={"report_type": "sales", "format": "pdf"},
        headers=admin_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_export_audit_and_notification(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    for report_type in ("audit", "notification"):
        response = await api_client.get(
            "/api/v1/reports/export",
            params={"report_type": report_type, "format": "xlsx"},
            headers=main_admin_headers,
        )
        assert response.status_code == 200
        assert len(response.content) > 100


@pytest.mark.asyncio
async def test_inventory_report_serial_filter(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    response = await api_client.get(
        "/api/v1/reports/inventory",
        params={"serial_number": "SN-RPT-001"},
        headers=main_admin_headers,
    )
    assert response.status_code == 200
    rows = response.json()["data"]["rows"]
    assert len(rows) == 1
    assert rows[0]["serial_number"] == "SN-RPT-001"


@pytest.mark.asyncio
async def test_inventory_report_date_filter(
    api_client: AsyncClient,
    db_session: AsyncSession,
    main_admin_headers: dict[str, str],
) -> None:
    await _seed_report_data(db_session)
    future = (datetime.now(UTC) + timedelta(days=1)).isoformat()
    response = await api_client.get(
        "/api/v1/reports/inventory",
        params={"date_from": future},
        headers=main_admin_headers,
    )
    assert response.status_code == 200
    assert response.json()["data"]["summary"]["total_rows"] == 0


@pytest.mark.asyncio
async def test_stream_inventory_uses_batches(db_session: AsyncSession) -> None:
    brand = await BrandRepository(db_session).create("BatchBrand")
    location = await LocationRepository(db_session).create(
        "BatchLoc",
        location_type=LocationType.WAREHOUSE,
    )
    product_model = await ProductModelRepository(db_session).create(
        brand_id=brand.id,
        model_number="BATCH-01",
        model_name="Batch Model",
        cpu="Intel i3",
        ram_gb=8,
        storage_value=Decimal("256"),
        storage_unit=StorageUnit.GB,
        storage_type=StorageType.SSD,
    )
    repo = InventoryItemRepository(db_session)
    for index in range(STREAM_BATCH_SIZE + 5):
        await repo.create(
            serial_number=f"SN-BATCH-{index:04d}",
            product_model_id=product_model.id,
            color="Black",
            current_location_id=location.id,
            status=InventoryStatus.AVAILABLE,
        )
    await db_session.commit()

    repo = ReportRepository(db_session)
    batch_sizes: list[int] = []
    with patch(
        "webstudio_backend.infrastructure.repositories.report_repository.STREAM_BATCH_SIZE",
        100,
    ):
        async for batch in repo.stream_inventory(ReportFilters(), batch_size=100):
            batch_sizes.append(len(batch))

    assert len(batch_sizes) >= 2
    assert sum(batch_sizes) == STREAM_BATCH_SIZE + 5


@pytest.mark.asyncio
async def test_report_service_aggregate_reports(db_session: AsyncSession) -> None:
    brand = await BrandRepository(db_session).create("AggBrand")
    location = await LocationRepository(db_session).create(
        "AggLoc",
        location_type=LocationType.WAREHOUSE,
    )
    product_model = await ProductModelRepository(db_session).create(
        brand_id=brand.id,
        model_number="AGG-01",
        model_name="Agg Model",
        cpu="Intel i7",
        ram_gb=32,
        storage_value=Decimal("1024"),
        storage_unit=StorageUnit.GB,
        storage_type=StorageType.SSD,
    )
    await InventoryItemRepository(db_session).create(
        serial_number="SN-AGG-001",
        product_model_id=product_model.id,
        color="Blue",
        current_location_id=location.id,
        status=InventoryStatus.AVAILABLE,
    )
    await db_session.commit()

    service = ReportService(db_session)
    location_rows = await service.location_report(ReportFilters())
    brand_rows = await service.brand_report(ReportFilters())
    model_rows = await service.product_model_report(ReportFilters())
    assert len(location_rows) == 1
    assert len(brand_rows) == 1
    assert len(model_rows) == 1
    assert location_rows[0].available == 1
